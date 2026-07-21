"""51_update_manuscript.py — apply corrected (insert-aware) results to the docx:
 - re-embed corrected Fig 14/15/16/18
 - rewrite paras 7,11,86,92,93,97,99,101,113,123 with corrected numbers/mutations
 - rebuild Table 5 (T4, engineering), Table 6 (T5, ddg), Table S2 (T7, hotspots)
"""
import pandas as pd
from docx import Document
from docx.shared import Inches
from docx.oxml.ns import qn
from PIL import Image

DOC = "manuscript/Manuscript_computational_draft.docx"
d = Document(DOC)


def set_text(i, txt):
    p = d.paragraphs[i]
    if p.runs:
        f = p.runs[0].font
        p.runs[0].text = txt
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(txt)


# ---------- text edits ----------
set_text(7,
    "These positions form a peripheral electrostatic (Glu/Lys) network distal to the "
    "active site whose charge state tracks temperature and stabilizes the fold "
    "cumulatively (FoldX −1.0 kcal/mol).")

set_text(11,
    "How enzymes gain thermal stability without sacrificing catalysis is a central "
    "question in protein evolution. Applying protein language models to a rigorously "
    "domain-verified dataset of glycoside hydrolases, we reveal the evolutionary logic "
    "of thermal adaptation in the (β/α)₈ TIM-barrel fold. ESM2 embeddings predict "
    "organism growth temperature and experimentally measured protein melting "
    "temperatures under leakage-controlled cross-validation, but their key contribution "
    "is mechanistic: they localize thermal adaptation of family GH5 to a sparse set of "
    "66 specific positions. These are not the catalytic residues — which are "
    "essentially invariant (alignment entropy 0.2 bits) and buried at the barrel axis "
    "— but evolutionarily permissive sites (entropy 3.1 bits) that lie distal to the "
    "active site, acquire charged (Glu/Lys) residues whose state is coupled to "
    "temperature, and stabilize the fold cumulatively when combined (up to −1.0 "
    "kcal/mol by FoldX). The same signature recurs in the TIM-barrel families GH1 and "
    "GH10 but not in the structurally distinct GH11 β-jelly-roll fold, defining a "
    "general principle: TIM-barrel enzymes hold a fixed catalytic core and tune "
    "thermostability through a sparse, peripheral electrostatic network at "
    "evolutionarily permissive positions. Independent directed-evolution experiments on "
    "a glycoside hydrolase — in which stabilizing mutations arise at peripheral, "
    "charge-altering, active-site-distal positions and act cumulatively — corroborate "
    "the mechanism. The work reframes machine learning from a black-box predictor into a "
    "means of discovering the rules of enzyme thermal evolution.")

set_text(86,
    "Two findings indicate that these positions constitute a genuine, GH5-specific "
    "adaptation mechanism rather than a restatement of bulk composition. First, the "
    "charge state of just these 66 positions — a sparse fraction of the fold — "
    "recapitulates optimal growth temperature almost as well as the full language model "
    "(Spearman = 0.80), and does so better than every one of 500 random sets of 66 "
    "positions (100th percentile; Figure 14c). Thermal adaptation in GH5 is thus "
    "concentrated at a definable set of positions, not spread diffusely across the "
    "sequence. Second, mapped onto a GH5 crystal structure (T. maritima, 3AMC) using an "
    "insert-aware alignment-to-structure mapping, none of the hotspots coincides with "
    "the two catalytic glutamates (Glu136/Glu253), and 53 of the 58 structurally "
    "mappable positions lie more than 12 Å from them (Figure 14b): charge is acquired "
    "across the periphery of the fold while the catalytic centre is left unperturbed. "
    "Indeed, the hotspots are on average significantly more distal from the catalytic "
    "glutamates than the remaining positions (mean 18.8 versus 16.1 Å; Mann–Whitney "
    "p < 0.001), consistent with a peripheral rather than active-site-proximal "
    "adaptation that is examined structurally in Section 3.16.")

set_text(92,
    "Mapping the charge hotspots onto the GH5 fold with an insert-aware alignment-to-"
    "structure correspondence and accurate solvent-accessibility calculations "
    "(FreeSASA) showed that they are both predominantly away from the catalytic centre "
    "(53 of the 58 structurally mappable positions > 12 Å from the catalytic "
    "glutamates) and significantly enriched at the solvent-exposed surface (66% "
    "surface-exposed; p < 0.001; Figure 15b,c). Charge is thus preferentially acquired "
    "at exposed peripheral positions of the barrel, where it can form solvent-facing "
    "salt bridges while sparing the buried active site — the hallmark of a peripheral "
    "electrostatic network. These surface, active-site-distal positions provide the "
    "safest targets for stabilizing engineering. From them we derived concrete, "
    "testable point mutations for a mesophilic GH5 (Bacillus sp., PDB 3PZT): at four "
    "surface, active-site-distal positions the mesophilic enzyme carries an uncharged "
    "residue where thermophiles overwhelmingly carry E or K (Table 5). The strongest "
    "family signal, A222K, lies at a position charged in 67% of thermophiles but only "
    "24% of mesophiles (Δ = 43 percentage points). These substitutions convert the "
    "discovery into experimentally verifiable hypotheses: each introduces a "
    "salt-bridge-forming residue at a surface position where thermophiles strongly "
    "favour one.")

set_text(93,
    "We assessed these mutations with two independent stability predictors — a "
    "protein-language-model zero-shot score (ESM-2 masked-marginal log-likelihood "
    "ratio) and physics-based ΔΔG calculations (FoldX, five runs on the repaired "
    "structure) — alongside the family signal (Table 6). The picture that emerged is "
    "instructive rather than a simple list of stabilizing hits. By FoldX, only N110E is "
    "clearly stabilizing (ΔΔG = −0.66 ± 0.14 kcal/mol); N244K and A222K are "
    "essentially neutral (−0.27 ± 0.02 and +0.15 ± 0.01 kcal/mol) and S248E is "
    "destabilizing (+0.75 ± 0.07 kcal/mol) despite its strong family signal, while the "
    "ESM-2 score — biased toward the natural mesophilic residue — disagreed with FoldX "
    "on several positions. Thus, although the hotspot positions robustly mark where GH5 "
    "lineages adapt, no single charge substitution transplanted into this mesophilic "
    "scaffold confers strong stabilization, and the three signals do not fully agree. "
    "This indicates that GH5 thermal adaptation is largely cumulative rather than "
    "reducible to individual dominant substitutions — a biologically meaningful "
    "conclusion that cautions against naive consensus engineering. The physics-based "
    "analysis nominates N110E as the most promising individual starting point, and the "
    "hotspot set as a whole defines the search space for combinatorial, "
    "experimentally guided stabilization. We provide the FoldX inputs and outputs with "
    "the deposited data.")

set_text(97,
    "The single-mutation results suggested that GH5 thermal adaptation is cumulative; we "
    "tested this directly with FoldX. Of the four surface, active-site-distal hotspot "
    "substitutions, the two most favourable (N110E −0.66 and N244K −0.27 kcal/mol) were "
    "combined and their effect compared with the sum of the single-mutation values "
    "(Figure 16). The double mutant N110E/N244K reached ΔΔG = −1.00 kcal/mol — "
    "substantially more stabilizing than either mutation alone and closely matching the "
    "additive expectation of −0.94 kcal/mol (coupling −0.07 kcal/mol), i.e. essentially "
    "additive with no significant epistasis. Extending the combination did not improve "
    "stability: adding the neutral A222K gave −0.74 kcal/mol, and further adding the "
    "destabilizing S248E collapsed the gain to −0.05 kcal/mol, confirming that not "
    "every family hotspot is a viable engineering site and that useful combinations "
    "must be assembled from individually compatible substitutions. The best combination "
    "therefore stabilizes the mesophilic scaffold by −1.00 kcal/mol, roughly 50% beyond "
    "the best single mutation. This directly confirms that GH5 thermal adaptation is "
    "cumulative: the position-resolved hotspot map is actionable not as isolated point "
    "mutations but as a set of compatible substitutions that together confer "
    "substantial predicted stabilization, providing a concrete multi-point engineering "
    "hypothesis for experimental test.")

set_text(99,
    "Figure 16. FoldX ΔΔG of engineered mutations in the mesophilic GH5 (3PZT). (a) "
    "Single surface-hotspot substitutions (blue, stabilizing; red, destabilizing). (b) "
    "Combinations: observed ΔΔG (blue) versus the additive sum of single effects "
    "(grey), showing cumulative, essentially additive stabilization reaching −1.00 "
    "kcal/mol for the N110E/N244K double mutant.")

# 3.18: 67 -> 66 (run-level, preserve formatting)
p101 = d.paragraphs[101]
for r in p101.runs:
    if "67" in r.text:
        r.text = r.text.replace("88, 67 and 80", "88, 66 and 80").replace(" 67 ", " 66 ")
if "88, 67" in p101.text:  # fallback whole-para
    set_text(101, p101.text.replace("88, 67 and 80", "88, 66 and 80"))

set_text(113,
    "The central result of this work is not a thermostability predictor but the "
    "evolutionary mechanism it reveals. By localizing thermal adaptation to specific "
    "positions and then asking why those positions, we arrive at a simple and general "
    "principle for the TIM-barrel fold: the catalytic core is held fixed while thermal "
    "stability is tuned through a sparse, peripheral electrostatic network. Three "
    "observations establish this. First, the catalytic residues are essentially "
    "invariant (alignment entropy 0.2 bits) and buried at the barrel axis — a geometry "
    "that cannot be altered without destroying activity, as confirmed experimentally by "
    "loss-of-function on active-site substitution in related glycoside hydrolases. "
    "Second, thermal adaptation instead concentrates at 66 evolutionarily permissive "
    "positions that are among the most variable in the fold (entropy 3.1 bits), lie "
    "distal to the catalytic centre, and acquire charged (predominantly Glu/Lys) "
    "residues whose state tracks growth temperature. Third, these positions act not "
    "individually but cumulatively: combining favourable substitutions yields additive "
    "stabilization (FoldX ΔΔG to −1.0 kcal/mol) that no single mutation approaches. "
    "Evolution therefore does not — and structurally cannot — remodel the active site "
    "to cope with heat; it displaces thermal selection onto tolerant peripheral "
    "positions where charge is energetically coupled to fold stability (Figure 18).")

set_text(123,
    "Figure 18. Evolutionary mechanism of thermal adaptation in the TIM-barrel fold. "
    "(a) The GH5 structure (Thermotoga maritima, 3AMC): the two catalytic glutamates "
    "(yellow) are buried at the barrel centre (active-site cleft shaded grey), whereas "
    "the 66 thermal-adaptation hotspots (spheres) are distributed over the periphery — "
    "Lys/Arg blue, Asp/Glu and other residues red. (b) The evolutionary logic: the "
    "high-constraint catalytic core (alignment entropy 0.2 bits) cannot be mutated, so "
    "thermal selection acts on low-constraint peripheral residues (entropy 3.1 bits), "
    "driving temperature-coupled accumulation of Glu/Lys into a peripheral salt-bridge "
    "network that stabilizes the fold while preserving catalysis. (c) A representative "
    "peripheral salt-bridge network formed by charged hotspots in the 3AMC structure "
    "(dashed edges, oppositely charged pairs < 13 Å; Asp/Glu red, Lys/Arg blue). (d) "
    "FoldX ΔΔG shows that a single mutation is insufficient (best −0.66 kcal/mol, "
    "N110E) whereas combining favourable hotspot substitutions stabilizes the fold "
    "cumulatively, reaching −1.00 kcal/mol for the N110E/N244K double mutant.")

print("text edits done")

# ---------- Table 5 (T4): engineering mutations ----------
eng = pd.read_csv("ai_training/analysis/engineering_mutations.csv")
t4 = d.tables[4]
rows4 = [
    (r["mutation"], f"{r['thermo_charged_%']:.0f}", f"{r['meso_charged_%']:.0f}",
     f"{r['delta_%']:.0f}", f"{r['RSA_3PZT']:.2f}", f"{r['dist_catGlu_A']:.0f}")
    for _, r in eng.iterrows()]
for ri, vals in enumerate(rows4, start=1):
    for ci, v in enumerate(vals):
        t4.rows[ri].cells[ci].paragraphs[0].runs[0].text = str(v)

# ---------- Table 6 (T5): ddg combined ----------
dd = pd.read_csv("ai_training/analysis/ddg_combined.csv")
t5 = d.tables[5]
for ri, (_, r) in enumerate(dd.iterrows(), start=1):
    esm = f"{r['ESM2_LLR']:.2f}".replace("-", "−")
    ddg = f"{r['FoldX_ddG']:+.2f} ± {r['FoldX_SD']:.2f}".replace("-", "−")
    vals = [r["mutation"], f"{r['family_delta_%']:.0f}", esm, ddg, r["FoldX_verdict"]]
    for ci, v in enumerate(vals):
        t5.rows[ri].cells[ci].paragraphs[0].runs[0].text = str(v)

# ---------- Table S2 (T7): top-25 hotspots by delta ----------
h = pd.read_csv("ai_training/analysis/hotspot_supptable.csv").sort_values("delta_%", ascending=False).head(25)
t7 = d.tables[7]
# ensure enough rows
need = len(h) + 1
while len(t7.rows) < need:
    t7.add_row()
def fmt(v, integer=False):
    if pd.isna(v):
        return "—"
    return f"{v:.0f}" if integer else str(v)
for ri, (_, r) in enumerate(h.iterrows(), start=1):
    res = "—" if pd.isna(r["res3AMC"]) else f"{r['res3AMC']:.0f}"
    vals = [res, r["thermo_pref_residue"], f"{r['thermo_charged_%']:.0f}",
            f"{r['meso_charged_%']:.0f}", f"{r['delta_%']:.0f}",
            fmt(r["location"]), fmt(r["active_site"])]
    cells = t7.rows[ri].cells
    for ci, v in enumerate(vals):
        c = cells[ci]
        if c.paragraphs[0].runs:
            c.paragraphs[0].runs[0].text = str(v)
            for extra in c.paragraphs[0].runs[1:]:
                extra.text = ""
        else:
            c.paragraphs[0].add_run(str(v))
print("tables updated: T4", len(rows4), "T5", len(dd), "T7", len(h))

# ---------- re-embed figures ----------
figs = {"rId30": "manuscript/figures/Fig14_discovery.png",
        "rId31": "manuscript/figures/Fig15_engineering.png",
        "rId24": "manuscript/figures/Fig16_combinations.png",
        "rId32": "manuscript/figures/Fig18_model.png"}
for rid, png in figs.items():
    part = d.part.related_parts[rid]
    with open(png, "rb") as f:
        part._blob = f.read()
# refit inline-shape extents to new aspect ratio (fixed 6.5" width)
for sh in d.inline_shapes:
    emb = sh._inline.xpath(".//a:blip/@r:embed")
    if emb and emb[0] in figs:
        im = Image.open(figs[emb[0]]); w, hh = im.size
        tw = Inches(6.5); sh.width = tw; sh.height = int(tw * hh / w)
print("figures re-embedded + refit")

d.save(DOC)
print("saved", DOC)
