"""54_export_data_workbook.py — consolidate all curated data + reported results into ONE Excel
workbook for reviewers (manuscript/Data_All_for_Review.xlsx).

Sheet 00_Index lists every sheet with a description, source file and dimensions. No new analysis:
each sheet is a verbatim copy of a committed CSV. Run from the repository root.
"""
import os
import pandas as pd
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
A = os.path.join(ROOT, "ai_training", "analysis")
D = os.path.join(ROOT, "ai_training")
OUT = os.path.join(ROOT, "manuscript", "Data_All_for_Review.xlsx")

# (source path, sheet name <=31 chars, description) — ordered for the reviewer
SHEETS = [
    # --- core dataset ---
    (f"{D}/gh5_topt_labeled.csv", "Dataset_757",
     "757 domain-verified GH5 sequences: accession, organism, sequence, OGT (TEMPURA) and enzyme Topt (BRENDA) labels"),
    (f"{A}/all757_predictions.csv", "Predictions_757",
     "Per-sequence predicted OGT and P(thermostable) for all 757 sequences (leakage-controlled CV)"),
    (f"{D}/gh5_all_verified.csv", "Domain_verification",
     "PF00150 HMMER domain-verification result (bitscore/e-value) for the 757 sequences"),
    # --- prediction / model evaluation ---
    (f"{A}/plm_benchmark.csv", "PLM_benchmark",
     "ESM2 model-size benchmark (8M-650M) for OGT regression"),
    (f"{A}/sota_comparison.csv", "SOTA_baselines",
     "Comparison against baseline thermostability predictors"),
    (f"{A}/ablation.csv", "Ablation",
     "Feature / representation ablation study"),
    (f"{A}/external_validation.csv", "External_validation",
     "External validation summary"),
    (f"{A}/meltome_benchmark.csv", "Meltome_benchmark",
     "External melting-temperature (Meltome Atlas / FLIP) benchmark"),
    (f"{A}/xai_importance.csv", "XAI_importance",
     "TreeSHAP + permutation feature importance (interpretable RF)"),
    (f"{A}/seq_feature_stats.csv", "Seq_feature_stats",
     "Compositional feature comparison thermophile vs mesophile (Mann-Whitney, Cliff's delta)"),
    # --- position-level hotspots / signature ---
    (f"{A}/thermal_hotspots.csv", "Thermal_hotspots",
     "66 charge-enriched hotspot positions (Fisher exact, Benjamini-Hochberg FDR < 0.05)"),
    (f"{A}/hotspot_supptable.csv", "Hotspot_table",
     "Hotspot supplementary table (per-position thermo/meso charge frequencies)"),
    (f"{A}/group_stats_fdr.csv", "Group_stats_FDR",
     "Residue-category enrichment statistics with FDR"),
    (f"{A}/signature_scores.csv", "Signature_scores",
     "Sparse charge-signature score per sequence vs OGT"),
    (f"{A}/constraint.csv", "Evol_constraint",
     "Per-position evolutionary constraint across the 272 match columns"),
    # --- structure / salt bridges ---
    (f"{A}/saltbridge_survey.csv", "Saltbridge_survey_82",
     "Salt-bridge and surface-charge survey across 82 experimentally solved GH5 structures"),
    (f"{A}/hotspot_saltbridges_3AMC.csv", "Saltbridges_3AMC",
     "Salt bridges among hotspot residues in the thermophile 3AMC"),
    (f"{A}/structure_surface_stats.csv", "Surface_stats",
     "Per-structure surface statistics"),
    (f"{A}/structural_summary.csv", "Structural_summary",
     "Structural mapping summary"),
    (f"{A}/catalytic_conservation.csv", "Catalytic_conservation",
     "Conservation of the catalytic glutamate pair"),
    (f"{A}/domain_completeness.csv", "Domain_completeness",
     "Per-sequence PF00150 domain coverage"),
    (f"{A}/active_site_integrity_all.csv", "Active_site_integrity",
     "Active-site integrity check across candidate sequences"),
    # --- stability (FoldX / zero-shot) + engineering ---
    (f"{A}/singles_ddg.csv", "FoldX_singles",
     "FoldX ddG for single hotspot substitutions"),
    (f"{A}/ddg_combined.csv", "FoldX_singles_combined",
     "FoldX single-mutation ddG (combined table with family signal)"),
    (f"{A}/combo_ddg.csv", "FoldX_combinations",
     "FoldX ddG for mutation combinations with epistatic coupling"),
    (f"{A}/combo_candidates.csv", "Combo_candidates",
     "Candidate mutation combinations"),
    (f"{A}/ddg_zeroshot.csv", "ESM_zeroshot_ddg",
     "ESM-2 zero-shot ddG estimates (alignment-free)"),
    (f"{A}/engineering_mutations.csv", "Engineering_mutations",
     "Selected engineering targets (manuscript Table 2)"),
    (f"{A}/engineering_candidates_all.csv", "Engineering_all",
     "All engineering candidate positions"),
    # --- candidates / discovery ---
    (f"{A}/top_candidates_validated.csv", "Top_candidates_validated",
     "Validated novel thermostable candidates (Table 1; P_thermostable > 0.5 filter)"),
    (f"{A}/top_candidates.csv", "Top_candidates_raw",
     "Top candidates before the validation filter"),
    (f"{A}/discovery_summary.csv", "Discovery_summary",
     "Biological-discovery summary counts"),
    (f"{A}/baek_5X3A_foldagnostic.csv", "Baek_5X3A_case",
     "Independent directed-evolution case (Baek et al., 5X3A) analysed fold-agnostically"),
    # --- cross-family generalization ---
    (f"{A}/generalization.csv", "Generalization_GH",
     "Cross-family generalization across GH1 / GH5 / GH10 / GH11"),
]


def main():
    index_rows = []
    with pd.ExcelWriter(OUT, engine="openpyxl") as xl:
        # placeholder index first (rewritten after we know row counts)
        pd.DataFrame({"": []}).to_excel(xl, sheet_name="00_Index", index=False)
        for path, sheet, desc in SHEETS:
            if not os.path.exists(path):
                print(f"  ! missing, skipped: {path}")
                continue
            df = pd.read_csv(path)
            df.to_excel(xl, sheet_name=sheet[:31], index=False)
            index_rows.append({"Sheet": sheet[:31], "Description": desc,
                               "Source file": os.path.relpath(path, ROOT),
                               "Rows": len(df), "Columns": df.shape[1]})
        idx = pd.DataFrame(index_rows, columns=["Sheet", "Description", "Source file", "Rows", "Columns"])
        idx.to_excel(xl, sheet_name="00_Index", index=False)

        # --- light formatting ---
        wb = xl.book
        hdr_fill = PatternFill("solid", fgColor="1B9E77")
        hdr_font = Font(bold=True, color="FFFFFF")
        for ws in wb.worksheets:
            ws.freeze_panes = "A2"
            for c in ws[1]:
                c.fill = hdr_fill
                c.font = hdr_font
                c.alignment = Alignment(vertical="center")
            for col in ws.columns:
                letter = get_column_letter(col[0].column)
                longest = max((len(str(c.value)) for c in col if c.value is not None), default=10)
                ws.column_dimensions[letter].width = min(max(longest + 2, 10), 60)
        # move index first
        wb.move_sheet("00_Index", -(wb.index(wb["00_Index"])))
    print(f"wrote {OUT}  ({len(index_rows)} data sheets + index)")


if __name__ == "__main__":
    main()
